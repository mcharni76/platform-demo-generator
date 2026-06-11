#!/usr/bin/env python3
"""
generate_scope_doc.py — Creates a scope/customization Excel workbook for platform demos.

Usage:
    pip install openpyxl  # only dependency
    python scripts/generate_scope_doc.py \
        --customer "Saudi Aramco" \
        --slug aramco \
        --industry energy \
        --country "Saudi Arabia" \
        --output ./docs/aramco_scope_v1.xlsx

Output: 6-tab Excel workbook ready for stakeholder review.
"""

import argparse
from datetime import date
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl required. Install with: pip install openpyxl")
    raise SystemExit(1)


# --- Styling ---
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="29B5E8", end_color="29B5E8", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="E8F4FD", end_color="E8F4FD", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)


# --- Industry Defaults ---
INDUSTRY_PRIORITIES = {
    "government": {"Policy Intelligence": "MUST", "Data Classification": "MUST", "Data Masking": "MUST", "Optimization": "OPTIONAL", "Pricing": "OPTIONAL"},
    "finance": {"Data Masking": "MUST", "Data Classification": "MUST", "ML & Predictive AI": "MUST", "Policy Intelligence": "MUST"},
    "healthcare": {"Data Masking": "MUST", "ML & Predictive AI": "MUST", "Data Quality": "MUST", "Data Classification": "MUST", "Optimization": "OPTIONAL"},
    "energy": {"Performance at Scale": "MUST", "Analytics Dashboards": "MUST", "Dynamic Tables": "MUST", "Data Quality": "MUST"},
    "telecom": {"ML & Predictive AI": "MUST", "Cortex AI (NLP)": "MUST", "Dynamic Tables": "MUST", "Analytics Dashboards": "MUST"},
    "retail": {"Analytics Dashboards": "MUST", "ML & Predictive AI": "MUST", "Cortex AI (NLP)": "HIGH", "Dynamic Tables": "MUST"},
    "education": {"Analytics Dashboards": "MUST", "ML & Predictive AI": "HIGH", "Time Travel": "HIGH"},
    "logistics": {"Performance at Scale": "MUST", "Dynamic Tables": "MUST", "Analytics Dashboards": "MUST"},
}

ENTITY_TEMPLATES = {
    "government": [("BENEFICIARIES", "CITIZENS"), ("ENROLLMENTS", "SERVICE_REQUESTS"), ("PROGRAMS", "GOVERNMENT_SERVICES"), ("TRAINERS", "OFFICERS"), ("SKILL_ASSESSMENTS", "INSPECTIONS"), ("FEEDBACK", "CITIZEN_FEEDBACK"), ("EVENTS", "PUBLIC_EVENTS"), ("STARTUPS", "LICENSED_BUSINESSES")],
    "finance": [("BENEFICIARIES", "CUSTOMERS"), ("ENROLLMENTS", "TRANSACTIONS"), ("PROGRAMS", "PRODUCTS"), ("TRAINERS", "ADVISORS"), ("SKILL_ASSESSMENTS", "RISK_ASSESSMENTS"), ("FEEDBACK", "CUSTOMER_FEEDBACK"), ("EVENTS", "COMPLIANCE_EVENTS"), ("STARTUPS", "FINTECHS")],
    "healthcare": [("BENEFICIARIES", "PATIENTS"), ("ENROLLMENTS", "APPOINTMENTS"), ("PROGRAMS", "CARE_PROGRAMS"), ("TRAINERS", "PHYSICIANS"), ("SKILL_ASSESSMENTS", "DIAGNOSES"), ("FEEDBACK", "PATIENT_FEEDBACK"), ("EVENTS", "CLINICAL_EVENTS"), ("STARTUPS", "HEALTH_STARTUPS")],
    "energy": [("BENEFICIARIES", "EMPLOYEES"), ("ENROLLMENTS", "PRODUCTION_RECORDS"), ("PROGRAMS", "ASSETS"), ("TRAINERS", "ENGINEERS"), ("SKILL_ASSESSMENTS", "INSPECTIONS"), ("FEEDBACK", "INCIDENT_REPORTS"), ("EVENTS", "MAINTENANCE_EVENTS"), ("STARTUPS", "VENDORS")],
    "telecom": [("BENEFICIARIES", "SUBSCRIBERS"), ("ENROLLMENTS", "CALL_RECORDS"), ("PROGRAMS", "PLANS"), ("TRAINERS", "AGENTS"), ("SKILL_ASSESSMENTS", "NETWORK_ASSESSMENTS"), ("FEEDBACK", "SUPPORT_TICKETS"), ("EVENTS", "NETWORK_EVENTS"), ("STARTUPS", "RESELLERS")],
    "retail": [("BENEFICIARIES", "CUSTOMERS"), ("ENROLLMENTS", "ORDERS"), ("PROGRAMS", "PRODUCTS"), ("TRAINERS", "STORE_MANAGERS"), ("SKILL_ASSESSMENTS", "PRODUCT_REVIEWS"), ("FEEDBACK", "CUSTOMER_FEEDBACK"), ("EVENTS", "PROMOTIONAL_EVENTS"), ("STARTUPS", "BRANDS")],
    "education": [("BENEFICIARIES", "STUDENTS"), ("ENROLLMENTS", "ENROLLMENTS"), ("PROGRAMS", "COURSES"), ("TRAINERS", "INSTRUCTORS"), ("SKILL_ASSESSMENTS", "ASSESSMENTS"), ("FEEDBACK", "STUDENT_FEEDBACK"), ("EVENTS", "CAMPUS_EVENTS"), ("STARTUPS", "ALUMNI_STARTUPS")],
    "logistics": [("BENEFICIARIES", "CUSTOMERS"), ("ENROLLMENTS", "SHIPMENTS"), ("PROGRAMS", "ROUTES"), ("TRAINERS", "DRIVERS"), ("SKILL_ASSESSMENTS", "DELIVERY_ASSESSMENTS"), ("FEEDBACK", "CUSTOMER_FEEDBACK"), ("EVENTS", "DELIVERY_EVENTS"), ("STARTUPS", "CARRIERS")],
}

REGULATORY_MAP = {
    "Saudi Arabia": "PDPL", "UAE": "PDPL", "Kuwait": "PDPL", "Qatar": "PDPL",
    "Bahrain": "PDPL", "Oman": "PDPL", "Jordan": "PDPL", "Egypt": "PDPL",
    "France": "GDPR", "Germany": "GDPR", "UK": "GDPR", "Italy": "GDPR",
    "Spain": "GDPR", "Netherlands": "GDPR",
    "USA": "SOX/CCPA",
}

# --- Template Pages ---
TEMPLATE_PAGES = [
    (1, "Architecture Overview", "Core", "MUST", "Medallion Architecture", "Static — always first"),
    (2, "Platform Overview", "Core", "MUST", "Unified Data Platform", ""),
    (3, "Performance at Scale", "Core", "MUST", "Elastic Compute", ""),
    (4, "Analytics Dashboards", "Core", "MUST", "Window Functions, H3 Geospatial", "7 tabs, MapLibre map"),
    (5, "Time Travel", "Core", "HIGH", "AT/BEFORE/UNDROP", "Step-by-step wizard"),
    (6, "Disaster Recovery", "Core", "HIGH", "CLONE/UNDROP", "Two-tab wizard"),
    (7, "Data Lineage", "Core", "HIGH", "OBJECT_DEPENDENCIES", "Node graph"),
    (8, "Data Quality", "Core", "HIGH", "Data Metric Functions", "3-tab layout"),
    (9, "ML & Predictive AI", "Advanced", "HIGH", "FORECAST/ANOMALY/CLASSIFICATION", "3 NCIM cards"),
    (10, "Cortex AI (NLP)", "Advanced", "MEDIUM", "SENTIMENT/SUMMARIZE/TRANSLATE", "5 cards, Arabic text"),
    (11, "Query Optimization", "Advanced", "MEDIUM", "Result Cache/ACCOUNT_USAGE", "Mock fallback"),
    (12, "Cost & Pricing", "Advanced", "MEDIUM", "Pay-per-Use", "Mock fallback"),
    (13, "Dynamic Tables", "Advanced", "HIGH", "Declarative Pipelines", ""),
    (14, "Policy Intelligence", "Advanced", "HIGH", "Cortex Search (RAG)", "Regulatory docs"),
    (15, "Data Masking", "Advanced", "HIGH", "Column-Level Security", "PII protection"),
    (16, "Data Classification", "Advanced", "HIGH", "SYSTEM$CLASSIFY", ""),
    (17, "Ask {Customer}", "Core", "MUST", "Cortex Analyst", "NL-to-SQL chatbot"),
]


def create_scope_doc(customer: str, slug: str, industry: str, country: str, output: str):
    wb = Workbook()
    today = date.today().isoformat()
    regulation = REGULATORY_MAP.get(country, "Generic")
    industry_lower = industry.lower()

    # --- Tab 1: Customer Context ---
    ws = wb.active
    ws.title = "Customer Context"
    ws.append(["Field", "Value"])
    style_header(ws)
    context_rows = [
        ("Customer Name", customer),
        ("Slug", slug),
        ("Industry", industry),
        ("Country", country),
        ("Language", "en+ar" if country in ["Saudi Arabia", "UAE", "Kuwait", "Qatar", "Bahrain", "Oman", "Jordan", "Egypt"] else "en"),
        ("Regulatory Context", regulation),
        ("Brand Color", ""),
        ("Website", ""),
        ("Snowflake Account", ""),
        ("Connection Name", f"{slug}-deploy"),
        ("Target Path", f"~/Projects/{slug.upper()}"),
        ("Date Created", today),
        ("Created By", ""),
        ("Stakeholders", ""),
    ]
    for field, value in context_rows:
        ws.append([field, value])
    auto_width(ws)

    # --- Tab 2: Page Requirements ---
    ws = wb.create_sheet("Page Requirements")
    ws.append(["#", "Page Name", "Category", "Include?", "Priority", "Snowflake Feature", "Custom Requirements", "Notes"])
    style_header(ws)

    priorities = INDUSTRY_PRIORITIES.get(industry_lower, {})
    for num, name, category, default_priority, feature, notes in TEMPLATE_PAGES:
        priority = priorities.get(name, default_priority)
        display_name = name.replace("{Customer}", customer)
        ws.append([num, display_name, category, "✅", priority, feature, "", notes])

    # Add empty custom rows
    for i in range(18, 24):
        ws.append([i, "", "Custom", "❓", "", "", "", ""])

    auto_width(ws)

    # --- Tab 3: Data Domain Mapping ---
    ws = wb.create_sheet("Data Domain Mapping")
    ws.append(["#", "Template Entity (MISK)", "Customer Entity", "Table Name", "Est. Rows", "PII Columns", "Notes"])
    style_header(ws)

    entities = ENTITY_TEMPLATES.get(industry_lower, ENTITY_TEMPLATES["government"])
    row_estimates = ["100K", "2M", "50", "2K", "8M", "50K", "500", "500"]
    for i, (template, customer_entity) in enumerate(entities):
        ws.append([i + 1, template, customer_entity, customer_entity, row_estimates[i] if i < len(row_estimates) else "", "", ""])

    # Add empty custom rows
    for i in range(len(entities) + 1, len(entities) + 5):
        ws.append([i, "", "", "", "", "", ""])

    auto_width(ws)

    # --- Tab 4: Scenario Flow ---
    ws = wb.create_sheet("Scenario Flow")
    ws.append(["#", "Demo Order", "Page", "Time (min)", "Lead Hook", "Snowflake Feature", "Audience"])
    style_header(ws)

    ws.append([1, "Opening", "Architecture Overview", 2, "Every node is live and clickable", "Full Platform", "Executive"])
    ws.append([2, "Core", "Platform Overview", 3, "Single source of truth", "Unified Platform", "All"])
    ws.append([3, "Core", "Performance at Scale", 3, "10M+ records in seconds", "Elastic Compute", "Technical"])
    # Pre-fill more based on industry lead scenario
    ws.append(["", "", "", "", "", "", ""])
    ws.append(["N", "Closing", f"Ask {customer}", 3, "Any question, plain language", "Cortex Analyst", "All"])

    auto_width(ws)

    # --- Tab 5: Timeline & Sessions ---
    ws = wb.create_sheet("Timeline & Sessions")
    ws.append(["Session", "Name", "Scope", "Target Date", "Status", "Dependencies", "Notes"])
    style_header(ws)

    sessions = [
        ("S0", "Planning", "Scope doc, PLAN.md, memory", "", "⬜ Pending", "Intake complete", ""),
        ("S0.5", "Data Generation", "Synthetic CSVs", "", "⬜ Pending", "S0 approved", ""),
        ("S1", "Infrastructure", "config + deploy SQL", "", "⬜ Pending", "S0 approved", ""),
        ("S2", "Backend", "FastAPI + Cortex Search", "", "⬜ Pending", "S1 complete", ""),
        ("S3", "Frontend Core", "8 core pages + shared", "", "⬜ Pending", "S2 complete", ""),
        ("S4", "Frontend Advanced", "9+ pages + semantic model", "", "⬜ Pending", "S3 complete", ""),
        ("S5", "Deploy + Demo Pack", "SPCS + DEMO_SCRIPT", "", "⬜ Pending", "S4 complete", ""),
        ("S6", "Custom (if needed)", "Custom pages", "", "⬜ Pending", "S5 complete", "Only if custom pages added"),
    ]
    for row in sessions:
        ws.append(list(row))

    auto_width(ws)

    # --- Tab 6: Sign-Off ---
    ws = wb.create_sheet("Sign-Off")
    ws.append(["Field", "Value"])
    style_header(ws)
    ws.append(["Document Version", "1.0"])
    ws.append(["Scope Frozen Date", ""])
    ws.append(["Approved By", ""])
    ws.append(["Approval Date", ""])
    ws.append(["", ""])
    ws.append(["CHANGE LOG", ""])
    ws.append([f"v1.0 — Initial scope", today])
    ws.append(["", ""])
    ws.append(["", ""])
    ws.append(["SCOPE AGREEMENT:", ""])
    ws.append(["", "Pages marked ✅ MUST will be built."])
    ws.append(["", "Pages marked ❓ will be discussed before S3/S4."])
    ws.append(["", "Custom pages add ~1 session each."])
    ws.append(["", "Scope changes after sign-off reset the Timeline."])
    ws.append(["", "Data domain mapping drives DDL, API, and UI labels."])

    auto_width(ws)

    # --- Save ---
    output_path = Path(output).expanduser()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    print(f"\n✅ Scope document created: {output_path}")
    print(f"   Tabs: Customer Context | Page Requirements | Data Domain Mapping | Scenario Flow | Timeline | Sign-Off")
    print(f"\n   Next: Review Tab 2, confirm entities in Tab 3, fill dates in Tab 5, get sign-off in Tab 6.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate platform demo scope document")
    parser.add_argument("--customer", required=True, help="Customer display name")
    parser.add_argument("--slug", required=True, help="Customer slug (lowercase)")
    parser.add_argument("--industry", required=True, help="Industry vertical")
    parser.add_argument("--country", required=True, help="Country")
    parser.add_argument("--output", required=True, help="Output .xlsx path")
    args = parser.parse_args()
    create_scope_doc(args.customer, args.slug, args.industry, args.country, args.output)
